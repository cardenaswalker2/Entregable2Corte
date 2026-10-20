import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_banner(doc, title, subtitle, subject, teacher):
    # Header box table
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "0F766E")
    set_cell_margins(cell, top=200, bottom=200, left=200, right=200)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p.add_run("FUNDACIÓN UNIVERSITARIA TECNOLÓGICO COMFENALCO\n")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(255, 255, 255)
    
    r_fac = p.add_run("FACULTAD DE INGENIERÍA • PROGRAMA DE INGENIERÍA DE SISTEMAS\nVII SEMESTRE — PROYECTO DE AULA (SEGUNDO SEGUIMIENTO)\n\n")
    r_fac.font.name = "Arial"
    r_fac.font.size = Pt(9.5)
    r_fac.font.color.rgb = RGBColor(204, 251, 241)
    
    r_title = p.add_run(f"{title}\n")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(255, 255, 255)
    
    r_sub = p.add_run(f"ENTREGABLE OFICIAL: {subject.upper()}\nDocente Evaluador: {teacher}\n")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(10.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(240, 253, 250)
    
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(12)
    p_meta.paragraph_format.space_after = Pt(12)
    p_meta.add_run("Autores / Investigadores: ").bold = True
    p_meta.add_run("Luis Cárdenas, Cristóbal Villamil, Jose Castillo, Daniel Gutiérrez.\n")
    p_meta.add_run("Fecha: ").bold = True
    p_meta.add_run("Octubre de 2026 • Cartagena de Indias D.T. y C.\n")
    p_meta.add_run("Repositorio Oficial: ").bold = True
    p_meta.add_run("https://github.com/cardenaswalker2/ClinicaApp.git")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_heading_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    for r in h.runs:
        r.font.name = "Arial"
        if level == 1:
            r.font.color.rgb = RGBColor(15, 118, 110)
            r.font.size = Pt(13)
        elif level == 2:
            r.font.color.rgb = RGBColor(30, 41, 59)
            r.font.size = Pt(11.5)
    return h

# ==========================================
# 1. GENERATE DWA WORD DOCUMENT
# ==========================================
def generate_dwa_doc():
    doc = Document()
    add_header_banner(
        doc,
        title="CLINICAAPP: PLATAFORMA INTEGRAL PARA LA GESTIÓN VETERINARIA",
        subtitle="Informe Técnico de Arquitectura Spring Boot 3.5, Agente Inteligente Nova Brain y Seguridad",
        subject="Desarrollo Web Avanzado",
        teacher="Ing. Fabio García Ramírez"
    )
    
    add_heading_styled(doc, "1. Introducción y Alcance en Desarrollo Web Avanzado", level=1)
    p = doc.add_paragraph("El presente documento consolida la entrega formal y auditable para la asignatura de Desarrollo Web Avanzado en el marco del proyecto de aula ClinicaApp. Se estructura una arquitectura modular y desacoplada en 5 capas fundamentada en Spring Boot 3.5 con Java 21 LTS, potenciada con el asistente conversacional inteligente Nova Brain integrado mediante APIs RESTful seguras con Modelos de Lenguaje Masivos (LLM Llama 3.3 70B vía Groq Cloud), respaldo por motor heurístico local (MOTOR_LOGICO), control de acceso estricto (RBAC de 6 roles), autenticación federada OAuth2/OIDC y reconocimiento biométrico facial 128D.")
    
    add_heading_styled(doc, "2. Diagrama de Arquitectura de la Aplicación (Spring Boot 3.5 + LLM)", level=1)
    p2 = doc.add_paragraph("En concordancia con los lineamientos de la guía y tomando como referencia la estructura del proyecto Educhat-notes, la arquitectura de ClinicaApp desacopla las responsabilidades en 5 capas cardinales:")
    
    table = doc.add_table(rows=6, cols=3)
    table.style = 'Table Grid'
    headers = ["Capa", "Componentes Principales en ClinicaApp", "Responsabilidad Técnica"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        set_cell_background(cell, "0F766E")
        for r in cell.paragraphs[0].runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9.5)
            
    rows_data = [
        ("Capa 1: Presentación / Frontend", "Thymeleaf UI, JavaScript reactivo (app.js, chat.js), Web Speech API (Voz STT/TTS), Bootstrap 5.3 adaptable.", "Renderizado de vistas del cliente, captura de audio del usuario, sintetización por voz y llamadas fetch asíncronas REST."),
        ("Capa 2: Controller (@RestController / @Controller)", "NovaAIController (/api/nova-brain/think), UsuarioController, ClinicaController, CitaController, OptimizacionController.", "Recepción de peticiones HTTP, validación de DTOs (NovaRequest), orquestación y respuesta JSON (NovaResponse)."),
        ("Capa 3: Service (@Service - Lógica de Negocio)", "NovaAIService, UsuarioServiceImpl, CitaService, OptimizacionServiceImpl, MaintenanceFilter, NetworkDeviceFilter.", "Inyección de dependencias, ensamblado de prompt contextual RAG, gestión de fallback MOTOR_LOGICO y biometría facial."),
        ("Capa 4: Repository (Spring Data MongoDB)", "UsuarioRepository, MascotaRepository, ClinicaRepository, CitaRepository, ServicioRepository, LogInferenciaRepository.", "Persistencia no relacional documental en MongoDB Atlas 7.0, consultas por Aggregation Framework e índices optimizados."),
        ("Capa 5: Proveedor LLM & Analítica Externa", "Groq Cloud API (Llama 3.3 70B Versatile vía ${GROQ_API_KEY}), Datasets CSV Kimball en docs/powerbi_datawarehouse/.", "Inferencia generativa ultra-rápida en la nube y repositorio de 3.469 hechos preparados para cuadros de mando.")
    ]
    
    for row_idx, data in enumerate(rows_data, start=1):
        for col_idx, text in enumerate(data):
            c = table.cell(row_idx, col_idx)
            c.text = text
            if row_idx % 2 == 0:
                set_cell_background(c, "F8FAFC")
            for r in c.paragraphs[0].runs:
                r.font.size = Pt(8.5)
                
    add_heading_styled(doc, "3. Prototipo de Clases del Agente Inteligente (Nova Brain)", level=1)
    doc.add_paragraph("Siguiendo el patrón arquitectónico del agente conversacional, se estructuran las siguientes clases:")
    doc.add_paragraph("• NovaAIController.java: Controlador REST que expone el endpoint /api/nova-brain/think. Extrae la identidad del usuario desde SecurityContextHolder, consulta en MongoDB sus mascotas registradas y citas pendientes, y construye el prompt del sistema.")
    doc.add_paragraph("• NovaRequest.java / NovaResponse.java: Objetos de transferencia de datos (DTOs) para tipar fuertemente las solicitudes, intenciones detectadas (TRIAGE_EMERGENCIA, INFO_SERVICIOS, AGENDAMIENTO) y acciones recomendadas.")
    doc.add_paragraph("• LogInferencia.java / LogInferenciaRepository.java: Entidad y repositorio para auditoría y observabilidad en MongoDB, guardando: usuarioId, prompt, respuesta, tokens consumidos y latencia en milisegundos.")
    doc.add_paragraph("• Motor de Contingencia (MOTOR_LOGICO): En caso de corte de conexión externa o falta de API key en producción, un analizador regex heurístico local responde inmediatamente garantizando disponibilidad 24/7.")
    
    add_heading_styled(doc, "4. Vistas Frontend e Integración con ChatBot / Asistente de Voz", level=1)
    doc.add_paragraph("• Vista asistente_voz.html: Widget moderno con Glassmorphism integrado en la interfaz de ClinicaApp.")
    doc.add_paragraph("• Reconocimiento de Voz (STT): Utilización de Web Speech API (webkitSpeechRecognition) permitiendo consultas habladas del usuario en lenguaje natural.")
    doc.add_paragraph("• Sintetizador de Voz (TTS): Respuesta hablada del agente inteligente mediante SpeechSynthesisUtterance.")
    doc.add_paragraph("• Validaciones Reactivas: Manejo visual de estados de carga, chips de preguntas rápidas y validación estricta contra envíos vacíos.")
    
    add_heading_styled(doc, "5. Seguridad Perimetral, Spring Security, RBAC y Biometría Facial", level=1)
    doc.add_paragraph("• SecurityConfig.java: Definición del SecurityFilterChain, protección contra ataques CSRF, y políticas de sesión.")
    doc.add_paragraph("• Control de Acceso Basado en Roles (RBAC): 6 roles estrictos: ROLE_ADMIN, ROLE_CLINICA, ROLE_VETERINARIO, ROLE_RECEPCIONISTA, ROLE_AUXILIAR y ROLE_CLIENTE.")
    doc.add_paragraph("• Autenticación Federada: Integración OAuth2 / OpenID Connect (OIDC) con Google y GitHub.")
    doc.add_paragraph("• Biometría Facial 128D: Algoritmo implementado en UsuarioServiceImpl.java que calcula la distancia euclidiana entre el vector facial registrado del usuario y la captura de cámara en vivo, validando el acceso con umbral estricto (< 0.38).")
    
    add_heading_styled(doc, "6. Modelo de Datos y Diccionario E-R (MongoDB Atlas 7.0)", level=1)
    doc.add_paragraph("Colecciones operacionales verificadas: usuarios (identidad y biometría), mascotas (perfil biológico), clinicas (georreferenciación y estado), citas (reservas médicas), servicios (catálogo de procedimientos), visitas (diagnóstico y facturación) y logs_inferencia (trazabilidad de IA).")
    
    doc.save("c:/Users/USUARIO/Music/CASA/EntregaProaula2Corte/Entregable_DWA_Fabio_Garcia_ClinicaApp.docx")

# ==========================================
# 2. GENERATE VD II WORD DOCUMENT
# ==========================================
def generate_vd_doc():
    doc = Document()
    add_header_banner(
        doc,
        title="CLINICAAPP: PLATAFORMA INTEGRAL PARA LA GESTIÓN VETERINARIA",
        subtitle="Informe Técnico de Identificación de Fuentes de Datos, Modelo Entidad-Relación (OLTP) y Modelo Dimensional Kimball (Star Schema)",
        subject="Visualización de Datos II",
        teacher="Ing. Walter Alberto Espriella Castellar"
    )
    
    add_heading_styled(doc, "1. Introducción y Justificación Analítica", level=1)
    doc.add_paragraph("En la gestión veterinaria asistencial, los registros clínicos y transaccionales permanecen convencionalmente confinados en sistemas transaccionales (OLTP) aislados. Esta condición impide a las direcciones médicas, coordinadores de sede y auditores administrativos disponer de una capa de inteligencia de negocios centralizada para la toma de decisiones oportunas sobre asistencia de pacientes, morbilidad de especies, tiempos de espera y rendimiento financiero.")
    doc.add_paragraph("El componente de Visualización de Datos II aborda esta necesidad formalizando la arquitectura de datos completa de ClinicaApp: en primer término, identificando rigurosamente las fuentes de datos operacionales; en segundo lugar, formalizando el Modelo Entidad-Relación (E-R) del sistema transaccional; y finalmente, estructurando la transición hacia el Modelo Dimensional en Esquema Estrella (Star Schema) bajo la metodología de Ralph Kimball, validado con un repositorio de 3.469 hechos y formulación de métricas DAX preparadas para su explotación analítica en Microsoft Power BI.")
    
    add_heading_styled(doc, "2. Identificación y Caracterización Rigurosa de Fuentes de Datos", level=1)
    doc.add_paragraph("Para responder a los requerimientos de la asignatura y garantizar la viabilidad analítica sin degradar el rendimiento del entorno transaccional, se identifican y clasifican las fuentes de información en dos estratos complementarios:")
    
    table = doc.add_table(rows=3, cols=4)
    table.style = 'Table Grid'
    headers = ["Fuente de Datos", "Tipo de Sistema", "Mecanismo de Acceso / Conexión", "Descripción y Entidades Auditadas"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        set_cell_background(cell, "0F766E")
        for r in cell.paragraphs[0].runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9)
            
    rows_data = [
        ("Base de Datos Operacional ClinicaApp (MongoDB Atlas 7.0 / Mapeo Relacional)", "OLTP (Transaccional Operacional)", "Spring Data MongoRepository / Conector ODBC/BI Connector / REST", "7 entidades transaccionales generadas en tiempo real: usuarios, mascotas, clínicas, citas, servicios, visitas y logs_inferencia."),
        ("Repositorio de Data Warehouse Sintético (docs/powerbi_datawarehouse/)", "Archivos Tabulares Planos (6 Datasets CSV UTF-8 Kimball)", "Conector nativo Texto/CSV en Power BI Desktop / Pandas Python", "Repositorio analítico preparado para BI (3.469 hechos de citas, 150 mascotas, 10 veterinarios, 6 sedes, 730 días) normalizado bajo la Ley 1581 de protección de datos.")
    ]
    for row_idx, data in enumerate(rows_data, start=1):
        for col_idx, text in enumerate(data):
            c = table.cell(row_idx, col_idx)
            c.text = text
            if row_idx % 2 == 0:
                set_cell_background(c, "F8FAFC")
            for r in c.paragraphs[0].runs:
                r.font.size = Pt(8.5)
                
    doc.add_paragraph()
    doc.add_paragraph("Nota Técnica sobre la Fuente Operacional: Si bien en el diseño conceptual relacional estándar se emplean motores RDBMS como MySQL o PostgreSQL, ClinicaApp implementa una arquitectura transaccional distribuida en MongoDB Atlas 7.0. No obstante, para efectos de modelado de datos y visualización analítica, la estructura de entidades, claves primarias (_id/PK) y claves foráneas (usuarioId, mascotaId, clinicaId, servicioId) preserva una estricta equivalencia relacional 1:N, garantizando total compatibilidad conceptual con los modelos relacionales clásicos.")

    add_heading_styled(doc, "3. Modelo Entidad-Relación (E-R) de la Base de Datos Transaccional (OLTP)", level=1)
    doc.add_paragraph("El Modelo Entidad-Relación captura las relaciones semánticas de la operación veterinaria diaria, asegurando integridad referencial, trazabilidad de citas médicas y auditoría de atenciones:")
    doc.add_paragraph("• Entidad USUARIOS: Almacena los actores del sistema (tutores, administradores, personal médico y de recepción). Clave primaria: id. Define roles bajo RBAC y almacena el vector facial biométrico de 128 dimensiones.")
    doc.add_paragraph("• Entidad MASCOTAS: Representa los pacientes animales. Clave foránea: usuario_id hacia USUARIOS (relación 1:N: un tutor puede tener múltiples mascotas). Almacena especie, raza, edad y patologías previas.")
    doc.add_paragraph("• Entidad CLINICAS: Registra las sedes físicas en Cartagena. Almacena georreferenciación GPS (latitud/longitud), capacidad operativa y estado administrativo.")
    doc.add_paragraph("• Entidad CITAS: Entidad transaccional central del sistema OLTP. Conecta con USUARIOS (tutor), MASCOTAS, CLINICAS y SERVICIOS mediante claves foráneas FK, tipando la fecha programada y el estado.")
    doc.add_paragraph("• Entidad SERVICIOS: Catálogo normalizado de procedimientos médicos, costos base y tiempos requeridos.")
    doc.add_paragraph("• Entidad VISITAS / ATENCIONES: Registro clínico post-atención que documenta tiempos de espera, duración del procedimiento médico, diagnóstico, monto facturado y evaluación del servicio.")
    
    # Insert E-R Diagram Image
    p_img_er = doc.add_paragraph()
    p_img_er.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_img_er = p_img_er.add_run()
    run_img_er.add_picture("c:/Users/USUARIO/Music/CASA/EntregaProaula2Corte/docs_diagrama_er_oltp.png", width=Inches(6.2))
    
    p_cap_er = doc.add_paragraph()
    p_cap_er.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap_er = p_cap_er.add_run("Figura 1: Diagrama Entidad-Relación (E-R) del Sistema Operacional Transaccional OLTP de ClinicaApp.")
    r_cap_er.font.size = Pt(8.5)
    r_cap_er.font.italic = True
    r_cap_er.font.color.rgb = RGBColor(100, 116, 139)

    add_heading_styled(doc, "4. Modelo Dimensional: Arquitectura en Esquema Estrella (Star Schema)", level=1)
    doc.add_paragraph("En entornos de analítica y Visualización de Datos (Power BI), los esquemas relacionales transaccionales 3NF resultan ineficientes debido a la proliferación de uniones (JOINs) costosas. Por este motivo, se adoptó la metodología de Ralph Kimball, transformando las entidades en un Esquema Estrella puro (Star Schema) compuesto por una Tabla de Hechos central rodeada de 5 Tablas Dimensionales conformes:")
    
    doc.add_paragraph("• Tabla de Hechos Central: Fact_Atenciones_Citas\n  - Grano: Una fila por cada cita individual atendida o programada (3.469 hechos auditados del periodo 2025-2026).\n  - Claves Foráneas (FK): id_fecha (INT), id_clinica (VARCHAR), id_servicio (VARCHAR), id_mascota (VARCHAR), id_veterinario (VARCHAR).\n  - Atributos Degenerados: canal_agendamiento (PORTAL_WEB, AGENTE_NOVA_CHATBOT, RECEPCION_PRESENCIAL, OPTIMIZACION_GEO_PLEB), estado_cita (COMPLETADA, CANCELADA, NO_ASISTIO), metodo_pago.\n  - Métricas Cuantitativas Aditivas: tiempo_espera_min, duracion_atencion_min, monto_bruto, monto_descuento, monto_neto, calificacion_satisfaccion (1-5), asistio (1/0), no_show (1/0), cancelada (1/0).")
    
    doc.add_paragraph("• Tablas Dimensionales Descriptivas (5 Dimensiones Conformes):\n  1. Dim_Tiempo.csv: 730 registros diarios (periodo 2025-2026). Columnas: id_fecha [PK], fecha_completa, anio, trimestre, mes_numero, mes_nombre, dia_mes, dia_semana, es_fin_semana.\n  2. Dim_Clinica.csv: 6 sedes en Cartagena. Columnas: id_clinica [PK], nombre, direccion, barrio, ciudad, telefono, estado_operativo, capacidad_diaria, latitud, longitud.\n  3. Dim_Servicio.csv: 8 servicios tipificados. Columnas: id_servicio [PK], nombre, categoria, precio_base, duracion_min, requiere_veterinario.\n  4. Dim_Mascota.csv: 150 animales. Columnas: id_mascota [PK], nombre, especie (Canino/Felino), raza, sexo, edad_anios, tamano, peso_kg, condicion_previa.\n  5. Dim_Veterinario.csv: 10 médicos. Columnas: id_veterinario [PK], nombre_completo, especialidad, experiencia_anios, id_clinica.")
    
    # Insert Star Schema Diagram Image
    p_img_star = doc.add_paragraph()
    p_img_star.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_img_star = p_img_star.add_run()
    run_img_star.add_picture("c:/Users/USUARIO/Music/CASA/EntregaProaula2Corte/docs_diagrama_estrella_kimball.png", width=Inches(6.2))
    
    p_cap_star = doc.add_paragraph()
    p_cap_star.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap_star = p_cap_star.add_run("Figura 2: Arquitectura del Modelo Dimensional en Esquema Estrella (Star Schema) de Ralph Kimball para Power BI.")
    r_cap_star.font.size = Pt(8.5)
    r_cap_star.font.italic = True
    r_cap_star.font.color.rgb = RGBColor(100, 116, 139)

    add_heading_styled(doc, "5. Métricas de Negocio y Fórmulas en Lenguaje DAX para Power BI", level=1)
    doc.add_paragraph("A partir del esquema estrella, se formalizaron las expresiones de análisis de datos (DAX) para la construcción de los cuadros de mando interactivos:")
    doc.add_paragraph("• Tasa de Asistencia Global:\n  Tasa_Asistencia = DIVIDE(COUNTROWS(FILTER(Fact_Atenciones_Citas, Fact_Atenciones_Citas[asistio] = 1)), COUNTROWS(Fact_Atenciones_Citas), 0)")
    doc.add_paragraph("• Tasa de Inasistencia (No-Show):\n  Tasa_NoShow = DIVIDE(COUNTROWS(FILTER(Fact_Atenciones_Citas, Fact_Atenciones_Citas[no_show] = 1)), COUNTROWS(Fact_Atenciones_Citas), 0)")
    doc.add_paragraph("• Ingresos Netos Facturados:\n  Ingresos_Netos = SUM(Fact_Atenciones_Citas[monto_neto])")
    doc.add_paragraph("• Adopción de Canales Digitales Autónomos:\n  Adopcion_Digital = DIVIDE(CALCULATE(COUNTROWS(Fact_Atenciones_Citas), Fact_Atenciones_Citas[canal_agendamiento] IN {\"AGENTE_NOVA_CHATBOT\", \"OPTIMIZACION_GEO_PLEB\"}), COUNTROWS(Fact_Atenciones_Citas), 0)")
    doc.add_paragraph("• Tiempo Promedio de Espera en Sala:\n  Tiempo_Promedio_Espera = AVERAGE(Fact_Atenciones_Citas[tiempo_espera_min])")
    
    add_heading_styled(doc, "6. Conclusiones y Nota de Rigor Académico sobre el Prototipo", level=1)
    p_nota = doc.add_paragraph("Nota de Rigor Académico y Cumplimiento Normativo: Se declara explícitamente que el modelo dimensional y las métricas DAX fueron validados técnica y funcionalmente mediante el repositorio sintético de 3.469 hechos. Este diseño desacopla deliberadamente la analítica del sistema transaccional OLTP de producción para evitar cuellos de botella por contención de bloqueos y proteger los datos personales de tutores y mascotas en cumplimiento estricto de la Ley Estatutaria 1581 de 2012 de Colombia, entregando una base semántica lista para ser cargada de manera inmediata en Power BI Desktop.")
    p_nota.paragraph_format.left_indent = Inches(0.25)
    
    doc.save("c:/Users/USUARIO/Music/CASA/EntregaProaula2Corte/Entregable_VD_Walter_Espriella_ClinicaApp.docx")


# ==========================================
# 3. GENERATE IO WORD DOCUMENT
# ==========================================
def generate_io_doc():
    doc = Document()
    add_header_banner(
        doc,
        title="CLINICAAPP: PLATAFORMA INTEGRAL PARA LA GESTIÓN VETERINARIA",
        subtitle="Informe Técnico de Optimización Geoespacial mediante Programación Lineal Entera Binaria (PLEB) y Haversine",
        subject="Investigación de Operaciones",
        teacher="Ing. Lucía Margarita Mestre Tapia"
    )
    
    add_heading_styled(doc, "1. Planteamiento del Problema Real en Cartagena de Indias", level=1)
    doc.add_paragraph("En situaciones de urgencia veterinaria o necesidad de consulta programada en el Distrito de Cartagena, los tutores de animales de compañía enfrentan desorientación espacial, saturación del tráfico urbano y desconocimiento de la operatividad de las sedes clínicas. Los propietarios realizan búsquedas manuales o empíricas que con frecuencia resultan en traslados a clínicas cerradas o innecesariamente lejanas, incrementando el tiempo de traslado y comprometiendo la estabilidad clínica del paciente animal.")
    doc.add_paragraph("Para resolver esta problemática, el componente de Investigación de Operaciones formaliza e implementa un modelo de Programación Lineal Entera Binaria (PLEB) acoplado a la ecuación geodésica del Semiverseno (Haversine), el cual garantiza la asignación óptima de la sede veterinaria factible de menor distancia geodésica al usuario dentro de su radio de desplazamiento.")
    
    add_heading_styled(doc, "2. Formulación Matemática Formal del Modelo PLEB", level=1)
    doc.add_paragraph("La formulación rigurosa del modelo de optimización se define a continuación:")
    
    doc.add_paragraph("2.1 Conjuntos e Índices:\n• I = {1, 2, ..., n}: Conjunto de centros veterinarios afiliados a la plataforma ClinicaApp en Cartagena.")
    
    doc.add_paragraph("2.2 Parámetros del Sistema:\n• (lat_u, lon_u): Coordenadas GPS de latitud y longitud del usuario / propietario de la mascota.\n• (lat_i, lon_i): Coordenadas GPS de latitud y longitud de la clínica veterinaria i ∈ I.\n• d_i: Distancia geodésica ortodrómica en kilómetros entre el usuario y la clínica i, calculada mediante la fórmula de Haversine (d_i ≥ 0).\n• A_i ∈ {0, 1}: Parámetro binario de operatividad institucional (1 si la clínica i está en estado APROBADA y activa, 0 si está PENDIENTE o inactiva).\n• D_max: Radio máximo de desplazamiento geográfico permitido en kilómetros (ej. 15.0 km).")
    
    doc.add_paragraph("2.3 Variable de Decisión:\n• X_i ∈ {0, 1}, para todo i ∈ I:\n  X_i = 1 si la clínica veterinaria i es seleccionada para atender al usuario.\n  X_i = 0 en caso contrario.")
    
    doc.add_paragraph("2.4 Función Objetivo:\nMinimizar la distancia geodésica total de traslado del paciente animal:\n  Min Z = ∑_{i ∈ I} (d_i · X_i)")
    
    doc.add_paragraph("2.5 Restricciones del Modelo:\na) Restricción de Selección Única (Asignación Exacta):\n   ∑_{i ∈ I} X_i = 1\n   (Garantiza que el algoritmo asigne exactamente una sede veterinaria al usuario).\n\nb) Restricción de Disponibilidad Operativa:\n   X_i ≤ A_i,  ∀ i ∈ I\n   (Impide asignar clínicas inactivas o en revisión administrativa; si A_i = 0, obliga X_i = 0).\n\nc) Restricción de Cobertura Geográfica Máxima:\n   d_i · X_i ≤ D_max,  ∀ i ∈ I\n   (Garantiza que la sede seleccionada no supere el radio límite de traslado establecido).\n\nd) Condiciones de Binariedad y No Negatividad:\n   X_i ∈ {0, 1},  ∀ i ∈ I\n   d_i ≥ 0,  ∀ i ∈ I")
    
    add_heading_styled(doc, "3. Formulación Geodésica del Semiverseno (Haversine)", level=1)
    doc.add_paragraph("Dado que las distancias planas Euclidianas introducen errores significativos sobre la curvatura de la superficie terrestre, se aplica la formulación trigonométrica de Haversine sobre coordenadas esféricas:")
    doc.add_paragraph("  d_i = 2 · R · arcsin( √( sin²(Δlat/2) + cos(lat_u) · cos(lat_i) · sin²(Δlon/2) ) )")
    doc.add_paragraph("Donde R = 6.371 km representa el radio medio volumétrico de la Tierra, Δlat = lat_i - lat_u y Δlon = lon_i - lon_u expresados en radianes.")
    
    add_heading_styled(doc, "4. Implementación en Código Fuente Java y Validación", level=1)
    doc.add_paragraph("El algoritmo de solución fue implementado en el backend Spring Boot dentro de la clase OptimizacionServiceImpl.java (interfaz IOptimizacionService). El servicio evalúa la factibilidad de cada clínica registrada, calcula d_i mediante Haversine y determina la variable binaria ganadora X_i = 1 que minimiza Z.")
    doc.add_paragraph("La funcionalidad se encuentra desplegada interactivamente en la interfaz web en la ruta /usuario/optimizacion-clinicas mediante mapas interactivos con Leaflet, permitiendo al usuario visualizar su posición, las sedes factibles y el trazado de la ruta óptima.")
    
    doc.save("c:/Users/USUARIO/Music/CASA/EntregaProaula2Corte/Entregable_IO_Lucia_Mestre_ClinicaApp.docx")

if __name__ == "__main__":
    generate_dwa_doc()
    generate_vd_doc()
    generate_io_doc()
    print("3 Word documents generated successfully!")
